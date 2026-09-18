"""Reusable python-pptx builder matching the established Saint Pierre
navy/gold lesson-deck visual system (colors and layout lifted directly
from the live rl7-3b-story-elements-lesson-map.pptx). 10in x 5.625in
slides throughout.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY = RGBColor(0x0F, 0x2E, 0x5C)
HEADER_BLUE = RGBColor(0x1B, 0x4E, 0x9B)
GOLD = RGBColor(0xFF, 0xC6, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SUBTITLE_BLUE = RGBColor(0xC9, 0xD9, 0xF0)
INK = RGBColor(0x1E, 0x2A, 0x3A)
GRAY = RGBColor(0x55, 0x63, 0x72)
BOX_BLUE = RGBColor(0xEE, 0xF4, 0xFC)
BOX_CREAM = RGBColor(0xFF, 0xF8, 0xE2)
W, H = Emu(9144000), Emu(5143500)

class Deck:
    def __init__(self, standard_code, standard_title):
        self.p = Presentation()
        self.p.slide_width = W
        self.p.slide_height = H
        self.blank = self.p.slide_layouts[6]
        self.code = standard_code
        self.title = standard_title
        self.n = 0

    def _slide(self):
        s = self.p.slides.add_slide(self.blank)
        self.n += 1
        return s

    def _rect(self, s, x, y, w, h, color, line=False):
        shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = color
        if line:
            shp.line.color.rgb = RGBColor(0xC7, 0xD0, 0xD6)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def _text(self, s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.15, space_after=6):
        """runs: list of paragraphs; each paragraph is a list of (text, size, bold, color) or a single tuple."""
        box = s.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        first = True
        for para_runs in runs:
            if isinstance(para_runs, tuple):
                para_runs = [para_runs]
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            para.alignment = align
            para.line_spacing = line_spacing
            para.space_after = Pt(space_after)
            for text, size, bold, color in para_runs:
                r = para.add_run()
                r.text = text
                r.font.size = Pt(size)
                r.font.bold = bold
                r.font.color.rgb = color
                r.font.name = 'Verdana'
        return box

    def _footer(self, s, kicker):
        self._text(s, 411480, 4709160, 4572000, 274320,
                    [[(self.code + '  ·  ' + self.title, 11, False, GRAY)]])
        self._text(s, 8412480, 4709160, 457200, 274320,
                    [[(str(self.n), 11, False, GRAY)]], align=PP_ALIGN.RIGHT)

    def cover(self, kicker, title, subtitle, tagline):
        s = self._slide()
        self._rect(s, 0, 0, 9144000, 5143500, NAVY)
        self._rect(s, 0, 0, 320040, 5143500, GOLD)
        # A title over ~30 characters wraps to two lines at this 36pt size;
        # push the rest of the layout down so the wrapped line never crowds the subtitle.
        extra = 411480 if len(title) > 30 else 0
        self._text(s, 822960, 1234440, 7772400, 274320, [[(kicker, 11, True, GOLD)]])
        self._text(s, 822960, 1645920, 7863840, 1005840 + extra, [[(title, 36, True, WHITE)]])
        self._text(s, 822960, 2743200 + extra, 7589520, 457200, [[(subtitle, 15, False, SUBTITLE_BLUE)]])
        self._rect(s, 822960, 3794760 + extra, 2926080, 45720, GOLD)
        self._text(s, 822960, 3931920 + extra, 7772400, 365760, [[(tagline, 11, False, WHITE)]])
        return s

    def header(self, s, kicker, title):
        self._rect(s, 0, 0, 9144000, 868680, HEADER_BLUE)
        self._rect(s, 0, 868680, 9144000, 64008, GOLD)
        self._text(s, 411480, 91440, 8321040, 228600, [[(kicker, 11, True, GOLD)]])
        self._text(s, 411480, 301752, 8321040, 502920, [[(title, 24, True, WHITE)]])

    def two_box(self, kicker, title, left_head, left_lines, right_head, right_lines):
        s = self._slide()
        self.header(s, kicker, title)
        self._rect(s, 411480, 1188720, 4114800, 3291840, BOX_BLUE)
        self._rect(s, 4709160, 1188720, 4023360, 3291840, BOX_CREAM)
        self._text(s, 685800, 1371600, 3657600, 320040, [[(left_head, 15, True, HEADER_BLUE)]])
        self._text(s, 685800, 1783080, 3611880, 2560320, [[(t, 12, False, INK)] for t in left_lines], space_after=14)
        self._text(s, 4983480, 1371600, 3657600, 320040, [[(right_head, 15, True, RGBColor(0xB0, 0x7A, 0x0A))]])
        self._text(s, 4983480, 1783080, 3520440, 2560320, [[(t, 12, False, INK)] for t in right_lines], space_after=14)
        self._footer(s, kicker)
        return s

    def definition(self, idx, total, term, definition, example, tip):
        s = self._slide()
        self.header(s, f'DEFINITION {idx} OF {total}', term)
        self._rect(s, 411480, 1188720, 8321040, 1188720, BOX_BLUE)
        self._text(s, 685800, 1310640, 7772400, 228600, [[('DEFINITION', 10, True, HEADER_BLUE)]])
        self._text(s, 685800, 1554480, 7772400, 731520, [[(definition, 13, False, INK)]])
        self._rect(s, 411480, 2469600, 8321040, 1188720, BOX_CREAM)
        self._text(s, 685800, 2591520, 7772400, 228600, [[('EXAMPLE', 10, True, RGBColor(0xB0, 0x7A, 0x0A))]])
        self._text(s, 685800, 2835360, 7772400, 731520, [[(example, 13, False, INK)]])
        self._text(s, 411480, 3840480, 8321040, 457200, [[('★  ' + tip, 12, True, HEADER_BLUE)]])
        self._footer(s, 'DEFINITION')
        return s

    def passage(self, kicker, title, body_paragraphs, note=None, size=12, space_after=12):
        s = self._slide()
        self.header(s, kicker, title)
        self._rect(s, 411480, 1188720, 8321040, 3291840 if not note else 2834640, BOX_BLUE)
        self._text(s, 685800, 1371600, 7772400, (3291840 if not note else 2834640) - 365760,
                    [[(p, size, False, INK)] for p in body_paragraphs], space_after=space_after)
        if note:
            self._rect(s, 411480, 4114800, 8321040, 365760, RGBColor(0xFD, 0xEC, 0xDC))
            self._text(s, 685800, 4160520, 7772400, 274320, [[(note, 11, True, RGBColor(0xB0, 0x4A, 0x0A))]])
        self._footer(s, kicker)
        return s

    def task_card(self, idx, total, title, focus, left_head, left_steps, right_head, right_body):
        s = self._slide()
        self.header(s, f'TASK CARD {idx} OF {total}', 'Task Card: ' + title)
        self._rect(s, 411480, 1188720, 8321040, 457200, HEADER_BLUE)
        self._text(s, 685800, 1310640, 7772400, 228600, [[('FOCUS: ' + focus, 13, True, WHITE)]])
        self._rect(s, 411480, 1783080, 4114800, 2697480, BOX_BLUE)
        self._rect(s, 4709160, 1783080, 4023360, 2697480, BOX_CREAM)
        self._text(s, 685800, 1920240, 3657600, 320040, [[(left_head, 14, True, HEADER_BLUE)]])
        self._text(s, 685800, 2286000, 3611880, 2103120, [[(t, 11.5, False, INK)] for t in left_steps], space_after=10)
        self._text(s, 4983480, 1920240, 3657600, 320040, [[(right_head, 14, True, RGBColor(0xB0, 0x7A, 0x0A))]])
        self._text(s, 4983480, 2286000, 3520440, 2103120, [[(right_body, 11.5, False, INK)]])
        self._footer(s, f'TASK CARD {idx} OF {total}')
        return s

    def open_questions(self, kicker, title, items, start=1, note=None):
        """items: list of {'q':...} dicts, open-response (no choices)."""
        s = self._slide()
        self.header(s, kicker, title)
        top = 1188720
        avail = 3291840 if not note else 2834640
        per_h = avail // len(items)
        for i, it in enumerate(items):
            qy = top + i * per_h
            self._text(s, 411480, qy, 8321040, per_h - 20000, [[(f'{start+i}. {it["q"]}', 12, False, INK)]])
        if note:
            self._rect(s, 411480, top + avail, 8321040, 365760, RGBColor(0xFD, 0xEC, 0xDC))
            self._text(s, 685800, top + avail + 45720, 7772400, 274320, [[(note, 11, True, RGBColor(0xB0, 0x4A, 0x0A))]])
        self._footer(s, kicker)
        return s

    def mc_questions(self, kicker, title, questions, start=1):
        """questions: list of (prompt, [choices A-D])"""
        s = self._slide()
        self.header(s, kicker, title)
        y = 1188720
        avail = 3520440
        per_h = avail // len(questions)
        for i, (prompt, choices) in enumerate(questions):
            qy = y + i * per_h
            self._text(s, 411480, qy, 8321040, 274320, [[(f'{start+i}. {prompt}', 12.5, True, INK)]])
            choice_txt = '   '.join(f'{chr(65+j)}. {c}' for j, c in enumerate(choices))
            self._text(s, 411480, qy + 274320, 8321040, per_h - 274320, [[(choice_txt, 11, False, GRAY)]])
        self._footer(s, kicker)
        return s

    def answer_key(self, kicker, title, answers):
        """answers: list of (label, answer_letter, reason)"""
        s = self._slide()
        self.header(s, kicker, title)
        y = 1188720
        avail = 3520440
        per_h = avail // len(answers)
        for i, (label, letter, reason) in enumerate(answers):
            qy = y + i * per_h
            self._text(s, 411480, qy, 8321040, 228600,
                        [[(f'{label}  —  Answer: {letter}', 12.5, True, HEADER_BLUE)]])
            self._text(s, 411480, qy + 228600, 8321040, per_h - 228600, [[(reason, 11, False, INK)]])
        self._footer(s, kicker)
        return s

    def open_answer_key(self, kicker, title, items, start=1):
        s = self._slide()
        self.header(s, kicker, title)
        y = 1188720
        avail = 3291840
        per_h = avail // len(items)
        for i, it in enumerate(items):
            qy = y + i * per_h
            self._text(s, 411480, qy, 8321040, per_h - 10000,
                        [[(f'{start+i}. ', 11.5, True, HEADER_BLUE), (it['answer'], 11, False, INK)]])
        self._footer(s, kicker)
        return s

    def guide_note(self, kicker, title, lines):
        s = self._slide()
        self.header(s, kicker, title)
        self._rect(s, 411480, 1188720, 8321040, 3291840, BOX_BLUE)
        self._text(s, 685800, 1371600, 7772400, 2926080, [[(t, 13, False, INK)] for t in lines], space_after=16)
        self._footer(s, kicker)
        return s

    def exit_ticket(self, left_head, left_lines, right_head, right_lines):
        s = self._slide()
        self.header(s, 'WRAP UP', 'Exit Ticket')
        self._rect(s, 411480, 1188720, 4114800, 3291840, BOX_CREAM)
        self._rect(s, 4709160, 1188720, 4023360, 3291840, BOX_BLUE)
        self._text(s, 685800, 1371600, 3657600, 320040, [[(left_head, 15, True, RGBColor(0xB0, 0x7A, 0x0A))]])
        self._text(s, 685800, 1783080, 3611880, 2560320, [[(t, 12, False, INK)] for t in left_lines], space_after=16)
        self._text(s, 4983480, 1371600, 3657600, 320040, [[(right_head, 15, True, HEADER_BLUE)]])
        self._text(s, 4983480, 1783080, 3520440, 2560320, [[(t, 12, False, INK)] for t in right_lines], space_after=14)
        self._footer(s, 'WRAP UP')
        return s

    def save(self, path):
        self.p.save(path)
        return self.n
